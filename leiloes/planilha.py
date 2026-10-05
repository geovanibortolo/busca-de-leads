"""Geração da planilha de leads a partir dos leilões coletados."""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUNAS = [
    ("Prioridade", 11),
    ("Dias até última praça", 10),
    ("UF", 5),
    ("Município", 20),
    ("Modalidade", 13),
    ("Área (ha)", 10),
    ("1ª praça", 16),
    ("Lance 1ª (R$)", 15),
    ("2ª praça", 16),
    ("Lance 2ª (R$)", 15),
    ("Avaliação (R$)", 15),
    ("Matrícula", 11),
    ("Cartório / Comarca", 24),
    ("CCIR (INCRA)", 18),
    ("Processo", 27),
    ("Devedor / executado (conferir)", 34),
    ("Origem do nome", 10),
    ("CPF/CNPJ citados", 22),
    ("Credor / exequente (conferir)", 30),
    ("Titular citado na matrícula", 26),
    ("Leiloeiro", 20),
    ("Situação", 11),
    ("Título", 40),
    ("Anúncio", 14),
    ("Edital (PDF)", 12),
    ("Matrícula (PDF)", 12),
    ("Incluído em", 12),
    ("Descrição", 60),
    ("Fonte", 16),
    ("Também anunciado em", 40),
]

_CORES = {
    "URGENTE": "F8CBAD",
    "ALTA": "FFE699",
    "MÉDIA": "E2EFDA",
}


def prioridade(dias):
    if dias is None:
        return "SEM DATA"
    if dias <= 7:
        return "URGENTE"
    if dias <= 21:
        return "ALTA"
    return "MÉDIA"


def _linha(l, agora):
    p1 = next((p for p in l.pracas if p[0] == 1), None)
    p2 = next((p for p in l.pracas if p[0] == 2), None)
    ultima = l.ultima_praca
    dias = (ultima[1].date() - agora.date()).days if ultima else None
    return [
        prioridade(dias),
        dias,
        l.uf,
        l.municipio,
        l.modalidade,
        l.area_ha,
        p1[1] if p1 else None,
        p1[2] if p1 else None,
        p2[1] if p2 else None,
        p2[2] if p2 else None,
        l.valor_avaliacao,
        l.matricula,
        l.cartorio,
        l.ccir,
        l.processos[0] if l.processos else "",
        l.devedor,
        l.devedor_fonte,
        ", ".join(l.documentos[:4]),
        l.credor,
        l.titular,
        l.leiloeiro,
        l.situacao,
        l.titulo,
        l.url,
        l.link_edital,
        l.link_matricula,
        l.data_inclusao,
        l.descricao[:2000],
        l.fonte,
        " | ".join(l.tambem_em),
    ]


def ativos(leiloes, agora=None):
    """Só leilões com alguma praça ainda por acontecer."""
    agora = agora or datetime.now()
    return [l for l in leiloes if l.ultima_praca and l.ultima_praca[1] >= agora]


def gerar(leiloes, caminho, agora=None):
    agora = agora or datetime.now()
    linhas = [_linha(l, agora) for l in leiloes]
    linhas.sort(key=lambda r: (r[1] is None, r[1] if r[1] is not None else 0))

    wb = Workbook()
    ws = wb.active
    ws.title = "Leilões rurais"
    ws.append([c[0] for c in COLUNAS])
    for i, (_, larg) in enumerate(COLUNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = larg
        cel = ws.cell(row=1, column=i)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="375623")
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32

    idx = {nome: i + 1 for i, (nome, _) in enumerate(COLUNAS)}
    for r in linhas:
        ws.append(r)
        n = ws.max_row
        cor = _CORES.get(r[0])
        if cor:
            ws.cell(row=n, column=1).fill = PatternFill("solid", fgColor=cor)
        for col in ("1ª praça", "2ª praça"):
            ws.cell(row=n, column=idx[col]).number_format = "DD/MM/YYYY HH:MM"
        ws.cell(row=n, column=idx["Incluído em"]).number_format = "DD/MM/YYYY"
        for col in ("Lance 1ª (R$)", "Lance 2ª (R$)", "Avaliação (R$)"):
            ws.cell(row=n, column=idx[col]).number_format = "#,##0.00"
        ws.cell(row=n, column=idx["Área (ha)"]).number_format = "#,##0.00"
        for col, rotulo in (("Anúncio", "abrir"), ("Edital (PDF)", "edital"), ("Matrícula (PDF)", "matrícula")):
            c = ws.cell(row=n, column=idx[col])
            if c.value:
                c.hyperlink = c.value
                c.value = rotulo
                c.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    # Aba de resumo por UF e modalidade
    res = wb.create_sheet("Resumo")
    res.append(["UF", "Modalidade", "Leilões", "Urgentes (≤7 dias)", "Soma lances 1ª praça (R$)", "Com devedor identificado"])
    grupos = {}
    for r in linhas:
        g = grupos.setdefault((r[2], r[4]), [0, 0, 0.0, 0])
        g[0] += 1
        g[1] += r[0] == "URGENTE"
        g[2] += r[7] or 0
        g[3] += bool(r[idx["Devedor / executado (conferir)"] - 1])
    for (uf, mod), (q, u, soma, dev) in sorted(grupos.items()):
        res.append([uf, mod, q, u, soma, dev])
        res.cell(row=res.max_row, column=5).number_format = "#,##0.00"
    for i, larg in enumerate((6, 30, 10, 18, 24, 22), 1):
        res.column_dimensions[get_column_letter(i)].width = larg
        res.cell(row=1, column=i).font = Font(bold=True)
    res.append([])
    res.append([f"Gerado em {agora:%d/%m/%Y %H:%M}. Colunas de partes são extraídas do texto do anúncio e precisam ser conferidas no edital/matrícula."])

    wb.save(caminho)
    return len(linhas)
